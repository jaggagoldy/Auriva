from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path
from datetime import date

OUT = Path('docs/product-readiness-audit')
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / 'auriva-independent-product-readiness-assessment-2026-07-23.docx'
MD = OUT / 'README.md'

INK = '241F1A'; PINE = '0E7466'; HONEY = 'E8A24C'; RED = 'C9584E'; AMBER = 'C77D24'; GREEN = '3F9D5A'; MUTED = '6F6A63'; PALE = 'F4EEE6'; BORDER = 'ECE3D6'; WHITE = 'FFFFFF'
W = 9360

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def borders(cell, color=BORDER):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), color); b.append(e)
    tcPr.append(b)

def cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn('w:' + side))
        if node is None:
            node = OxmlElement('w:' + side); mar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def set_cell_width(cell, width):
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(width)); tcW.set(qn('w:type'), 'dxa')

def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'), 'dxa')
    tblInd = tblPr.first_child_found_in('w:tblInd')
    if tblInd is None:
        tblInd = OxmlElement('w:tblInd'); tblPr.append(tblInd)
    tblInd.set(qn('w:w'), '120'); tblInd.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(width)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i]); cell_margin(cell); borders(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def mark_header_row(row):
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement('w:tblHeader'); node.set(qn('w:val'), 'true'); trPr.append(node)

def font(run, size=10, bold=False, color=INK, italic=False):
    run.font.name = 'Aptos'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size); run.bold = bold; run.italic = italic; run.font.color.rgb = RGBColor.from_string(color)

def style_doc(d):
    sec = d.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.85); sec.left_margin = sec.right_margin = Inches(0.8)
    sec.header_distance = Inches(.35); sec.footer_distance = Inches(.35)
    normal = d.styles['Normal']; normal.font.name = 'Aptos'; normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos'); normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in [('Title', 28, PINE, 0, 6), ('Heading 1', 17, PINE, 16, 7), ('Heading 2', 13, PINE, 12, 5), ('Heading 3', 11, '083F37', 8, 3)]:
        s = d.styles[name]; s.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'; s._element.rPr.rFonts.set(qn('w:ascii'), s.font.name); s._element.rPr.rFonts.set(qn('w:hAnsi'), s.font.name); s.font.size = Pt(size); s.font.color.rgb = RGBColor.from_string(color); s.font.bold = name != 'Title'; s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
    header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run('AURIVA | INDEPENDENT PRODUCT READINESS ASSESSMENT | CONFIDENTIAL'); font(r, 7.5, True, MUTED)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run('Independent Product Review Board • 23 July 2026 • Page '); font(r, 8, False, MUTED)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def para(d, text='', bold_lead=None, keep=False):
    p = d.add_paragraph(); p.paragraph_format.keep_with_next = keep
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); font(r, 10, True)
        r = p.add_run(text[len(bold_lead):]); font(r, 10)
    else:
        r = p.add_run(text); font(r, 10)
    return p

def bullet(d, text):
    p = d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); font(r, 10); return p

def h(d, text, level=1): d.add_heading(text, level=level)

def table(d, headers, rows, widths, font_size=8.3):
    t = d.add_table(rows=1, cols=len(headers)); set_table_geometry(t, widths)
    mark_header_row(t.rows[0])
    for i, label in enumerate(headers):
        c = t.rows[0].cells[i]; shade(c, PINE); p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run(label); font(r, font_size, True, WHITE)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i];
            if len(t.rows) % 2 == 0: shade(c, 'FBF8F4')
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(val)); font(r, font_size)
    set_table_geometry(t, widths)
    d.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def callout(d, label, text, color='E4EFEC'):
    t = d.add_table(rows=1, cols=1); set_table_geometry(t, [W]); c=t.cell(0,0); shade(c,color)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(label + '  '); font(r,10,True,PINE if color != 'FBE5E3' else RED); r=p.add_run(text); font(r,10)
    d.add_paragraph().paragraph_format.space_after=Pt(1)

def pagebreak(d): d.add_page_break()

scores = [
    ('Overall Product Score', '5.2', 'RED', 'A credible encounter-to-cash foundation, not a commercially complete Healthcare OS.'),
    ('Market Readiness', '4.0', 'RED', 'Category-table stakes—messaging, online payments, intake/consent, reporting and mature scheduling—remain absent or thin.'),
    ('UI', '6.8', 'AMBER', 'Tokenized and intentional, but the assessment is source-led because the live UI was not reachable from the audit browser.'),
    ('UX', '5.4', 'AMBER', 'Happy-path stage ownership is strong; cross-workflow recovery, discovery and automation are incomplete.'),
    ('Engineering Maturity', '5.8', 'AMBER', 'Strong boundaries, audit/event seams and 74 test files; quality gates currently fail and scaling assumptions are explicit.'),
    ('Business Readiness', '3.7', 'RED', 'No finalized pricing, package design, proof of ROI, migration plan or support operating model.'),
    ('Launch Readiness', '3.5', 'RED', 'Production OTP, backups/restore, monitoring and release gates remain unchecked; lint fails.'),
    ('Commercial Readiness', '3.8', 'RED', 'Value proposition has promise but lacks a sellable outcome package and hard switching proof.'),
    ('Customer Delight', '4.7', 'RED', 'The warm patient shell is promising, but it does not communicate, collect payment, or guide care outside the clinic.'),
    ('Clinical Readiness', '4.8', 'RED', 'Useful documentation is present, yet structured clinical governance, consent, prescribing safeguards and longitudinal record depth are immature.'),
]

recs = [
('R-001','Critical','Launch','Live OTP delivery and delivery observability are an unchecked external-beta gate.','Patients cannot reliably log in; every patient-facing promise collapses.','Integrate one production SMS/OTP vendor, delivery callbacks, retries, rate limits and a daily failure dashboard.','Makes patient activation measurable and protects trust.','M','P0'),
('R-002','Critical','Launch','Backups and a production restore rehearsal are not evidenced as complete.','A recoverability failure with health data is existential.','Provision encrypted backups, define RPO/RTO, execute a documented restore drill, and make it a release gate.','Prevents catastrophic loss and reassures institutional buyers.','M','P0'),
('R-003','Critical','Engineering','The current lint command fails with 43 errors and 51 warnings.','A failing quality gate normalizes unsafe change and undermines release claims.','Fix all source lint errors; exclude archival prototypes deliberately; block merge and release on clean lint, types and tests.','Cuts regression cost and restores credibility.','M','P0'),
('R-004','Critical','Clinical safety','No documented clinical governance, sign-off, or specialty validation process is visible.','A clinically unsafe workflow can create patient harm and liability.','Create a Clinical Safety Board, risk classification, change-control, usability validation and release sign-off for clinical changes.','Reduces clinical and legal risk; becomes an enterprise sales asset.','M','P0'),
('R-005','Critical','Patient engagement','No outbound SMS, WhatsApp, email or push delivery exists; in-app alerts require patients to open Auriva.','No-show reduction, follow-up completion and document delivery cannot happen.','Ship consented multi-channel communications with templates, preferences, delivery status and cost controls before broad launch.','Unlocks retention, attendance and visible daily value.','L','P0'),
('R-006','Critical','QA','There is no evidenced browser end-to-end release suite for the full care day.','Service tests do not prove the real clinic can move a patient through the UI.','Automate role-based golden paths: setup, booking, walk-in, queue, consult, invoice, payment, portal and permission denial.','Reduces launch-day workflow breaks.','L','P0'),
('R-007','Critical','Security','In-memory rate limiting and reminder sweep are single-instance only.','Scale can weaken abuse protection and duplicate operational actions.','Move rate limits, idempotency and scheduled delivery locks to shared infrastructure before multi-instance deployment.','Prevents security and reliability incidents during growth.','M','P0'),
('R-008','Critical','Commercial','Pricing is explicitly not finalized.','Sales cannot package, quote or forecast a product without a defensible commercial offer.','Define editions, limits, add-ons, implementation fees, usage-metered messaging and a price-testing plan.','Enables selling, forecasting and controlled expansion.','M','P0'),
('R-009','Critical','Patient experience','Patients cannot pay online or from the portal.','The financial journey stops at reception; self-service and collection efficiency are lost.','Launch payment links/UPI intent/QR with receipts, reconciliation and failure handling.','Accelerates collection and makes the portal useful.','L','P0'),
('R-010','Critical','Findability','Reception and doctor surfaces lack a global patient search; lists have no pagination.','Busy staff lose time and risk opening the wrong patient as data grows.','Implement tenant-scoped global search by name, phone, health ID and DOB; add paging, server filtering and audit context.','Cuts handling time and improves safety.','L','P0'),
('R-011','High','Onboarding','Clinic setup lacks proven time-to-first-value instrumentation and guided data migration.','A buyer may never reach a successful first day.','Create an activation checklist, sample data, import tools, onboarding telemetry and concierge playbook.','Improves conversion and lowers support cost.','M','P1'),
('R-012','High','Scheduling','Calendar is day-only on reception; advanced recurrence, drag/drop and scheduling requests are deferred.','Auriva loses to commodity scheduling tools before its operations story is experienced.','Prioritize recurring schedules, reschedule, conflict explanations, cancellation policy and waitlist before broad rollout.','Reduces staff workload and no-shows.','L','P1'),
('R-013','High','Reception','Doctor assignment is fixed at creation; the board has no reassignment or queue transfer.','One late doctor creates avoidable queue congestion and poor patient experience.','Add a guarded reassign/transfer action with reason, patient notification and audit record.','Improves throughput without adding headcount.','M','P1'),
('R-014','High','Revenue','Consult services/treatments are not fully wired into a reliable consult-to-invoice commercial flow.','Revenue capture is shallow and dependent on manual work.','Make ordered services the canonical invoice input; show expected charge, waiver and exception reasons at handoff.','Raises capture rate and billing confidence.','L','P1'),
('R-015','High','Clinical','Clinical documentation remains largely free text on the Appointment “god table.”','Longitudinal care, safe analytics and interoperability are constrained.','Define a versioned clinical encounter model with structured diagnoses, vitals, problems, medications and provenance.','Creates an enduring clinical data asset.','XL','P1'),
('R-016','High','Prescription','Prescription safety is incomplete: no interaction checks, structured medication governance or pharmacy workflow.','Medication errors and support burden remain material.','Ship structured Rx documents, allergy/duplicate therapy checks, immutable signing, amendment trail and local formulary strategy.','Raises clinical trust and differentiates doctor workflow.','L','P1'),
('R-017','High','Consent','Digital intake, consent and signature workflows are absent.','Clinical and medico-legal documentation starts too late and remains weak.','Introduce configurable forms, versioned consent, e-signature, completion status and pre-visit gating.','Cuts front-desk work and improves defensibility.','L','P1'),
('R-018','High','Analytics','Command Center is an activity view, not a decision system; analytics are explicitly thin.','Owners cannot see whether Auriva improves cash, flow or retention.','Ship a small KPI pack: no-show, waiting time, consult completion, unpaid balance, rebooking and plan conversion.','Creates owner stickiness and renewal evidence.','M','P1'),
('R-019','High','Operations','Queue wait aging is visual only; capacity thresholds and intervention workflows are deferred.','The product displays congestion without helping resolve it.','Set configurable service-level thresholds, escalation states and one-action recovery options.','Protects patient experience and clinic throughput.','M','P1'),
('R-020','High','Data protection','No formal compliance posture, privacy mapping or data-processing pack is evidenced.','Institutional buyers will block procurement; trust claims are unsubstantiated.','Produce data-flow map, retention/deletion policy, breach response, vendor register, DPA and jurisdiction-specific legal review.','Unlocks procurement and reduces legal exposure.','M','P1'),
]

more = [
('R-021','High','Product','Treatment planning and procedure tracking are roadmap items, not current value.','Physio and dental customers cannot run their core multi-visit care model.','Complete plan → session → checkout → next-booking loop before targeting these specialties.','Opens high-LTV specialty segments.','L','P1'),
('R-022','High','Product','Clinical timeline is not yet the decisive pre-consult context screen.','Doctors still assemble history across surfaces.','Unify visits, documents, plans, labs, payments and active problems with role-specific views.','Reduces consult prep time.','L','P1'),
('R-023','High','UX','“Soon” and reserved features appear in several live role surfaces.','They signal incompleteness at the exact moment a buyer evaluates value.','Hide non-sellable navigation by entitlement; replace placeholders with meaningful next actions.','Improves perceived quality and task focus.','S','P1'),
('R-024','High','Support','No mature customer-success operating model or implementation SLA is evidenced.','First clinics will experience adoption risk as product failure.','Define pilot success plan, escalation paths, response targets, weekly value review and feedback intake.','Improves retention and referenceability.','M','P1'),
('R-025','High','Migration','No legacy data import and reconciliation journey is visible.','Switching cost remains too high for established clinics.','Provide patient, appointment, balance and document imports with dry run, mapping and exception report.','Removes a primary purchase barrier.','L','P1'),
('R-026','High','Payments','No documented tax/GST-ready commercial receipt strategy is evidenced; tax fields are inactive.','Indian clinics may not be able to use billing as their system of record.','Validate target segments; implement tax configuration, compliant invoice fields and accountant export where required.','Protects billing adoption.','M','P1'),
('R-027','High','Roles','Nurse and technician permissions exist but dedicated action surfaces are deferred.','Role promise is broader than usable work; teams revert to shared credentials or paper.','Build minimum safe vitals capture and results-entry workflows or remove roles from packaging.','Improves team adoption and access safety.','M','P1'),
('R-028','High','Market','Target segment is too broad: solo doctors through polyclinics, across many specialties.','Message, workflow and sales motion become generic.','Choose one beachhead (initially cash/UPI, multi-visit physio or dental) and say no to the rest.','Concentrates learning and GTM efficiency.','S','P1'),
('R-029','High','Market','Mental-health and nutrition workflows have no explicit fit validation.','Competing against SimplePractice/Healthie without forms, messaging or telehealth is unwinnable.','Avoid these segments until their must-have journey is built and validated.','Prevents expensive failed sales.','S','P1'),
('R-030','High','Marketing','The product claims “Healthcare OS” before sufficient operational breadth is delivered.','This creates expectation debt and a credibility gap in demos.','Sell “front-desk-to-cash operating system for [beachhead] clinics” until the OS claims are earned.','Raises conversion and trust.','S','P1'),
('R-031','Medium','Accessibility','Accessibility is implemented through scattered ARIA/focus patterns, with no automated audit evidence.','Keyboard and assistive-technology failures will surface late.','Add axe-based CI, keyboard test scripts, focus-order acceptance criteria and screen-reader smoke tests.','Reduces exclusion and enterprise risk.','M','P2'),
('R-032','Medium','Accessibility','Touch targets are often compact 28–36px controls in clinically busy areas.','Tablet and gloved/fast-use interaction errors increase.','Adopt a 44px minimum for primary touch actions and evaluate board density separately.','Fewer operational mistakes.','M','P2'),
('R-033','Medium','Accessibility','Dark-mode semantic colors are declared but not contrast-verified across component states.','Clinical status may become ambiguous in night use.','Run WCAG contrast tests for all status foreground/background combinations and focus states.','Improves safety and compliance.','S','P2'),
('R-034','Medium','UX','Error states are designed, but recovery ownership and retry semantics are inconsistent by workflow.','Users may retry unsafe mutations or abandon work.','Standardize mutation states: pending, success, failure, retry-safe, and support reference ID.','Reduces support and duplicate actions.','M','P2'),
('R-035','Medium','UX','Autosave in clinical forms has no evidenced visible save confidence or conflict resolution.','Clinicians cannot know whether a critical note persisted.','Show save status, last saved time, unsaved-warning and concurrent-edit behavior.','Builds clinician trust.','M','P2'),
('R-036','Medium','UX','The doctor “Pause booking” action is informational only.','A false control damages trust during a stressful day.','Remove it until real, or implement a scoped availability/booking hold with confirmation.','Eliminates expectation mismatch.','S','P2'),
('R-037','Medium','Search','Health ID search exists in the backend but is not surfaced in core UI.','A designed safety identifier is not delivering operational value.','Expose Health ID and DOB-assisted disambiguation in reception and doctor search.','Reduces wrong-patient risk.','S','P2'),
('R-038','Medium','Search','Billing and lab worklists lack patient/invoice/test search.','Staff must scan queues manually under pressure.','Add server-side search, active filters and persistent query state.','Cuts processing time.','M','P2'),
('R-039','Medium','Scheduling','Reception board and calendar are split rather than one operational workspace.','Staff context-switches between “who is here” and “what is next.”','Provide a linked day workspace with board/capacity/calendar panes and preserved filter context.','Improves front-desk speed.','L','P2'),
('R-040','Medium','Patient','Patient records cannot be searched.','Longitudinal records become unusable as history grows.','Add secure full-text or metadata search with document type/date filters.','Improves patient self-service.','L','P2'),
('R-041','Medium','Patient','Family sharing lacks granular consent and proxy governance.','Family convenience may cause privacy breaches.','Add per-profile consent, guardian role, expiration and audit trail before marketing family care broadly.','Protects trust and compliance.','L','P2'),
('R-042','Medium','Patient','The patient app shows medicines informationally with no adherence, refill or escalation loop.','The portal is passive after the visit.','Start with refill due and follow-up reminders, then validate adherence features by specialty.','Creates repeat engagement.','M','P2'),
('R-043','Medium','Patient','There is no “after the visit” closing moment.','Patients leave without a clear plan, documents or next action.','Create a concise care summary with next appointment, Rx, payment status and help channel.','Raises trust and rebooking.','M','P2'),
('R-044','Medium','Clinical','Lab result display is basic and “recent labs” is coming soon in doctor context.','Clinical information is incomplete at point of care.','Make results reviewable, trendable, attributable and acknowledged before claiming a lab workflow.','Improves safety and clinician value.','L','P2'),
('R-045','Medium','Clinical','No specialty template strategy is evidenced beyond generic SOAP.','Generic documentation produces either low adoption or unsafe custom workarounds.','Develop validated specialty starter packs only for the selected beachhead, with governance.','Improves adoption and data quality.','M','P2'),
('R-046','Medium','Clinical','No clear correction/amendment UX for signed clinical artifacts is evidenced.','Clinical records may be silently changed or hard to correct.','Adopt lock/sign/amend pattern with reason, timestamp, author and version view.','Protects medico-legal integrity.','M','P2'),
('R-047','Medium','Clinical','No allergies/contraindications data-quality workflow is evidenced.','Safety chips can be stale or missing.','Add structured capture, source, last-reviewed date and high-risk acknowledgment.','Improves meaningful safety support.','M','P2'),
('R-048','Medium','Design system','Semantic design tokens exist but brand color tokenization is documented as debt.','Visual changes risk inconsistent hard-coded implementation.','Finish semantic token inventory, component variants and visual regression baselines.','Speeds safe product evolution.','M','P2'),
('R-049','Medium','Design system','Two header conventions are intentionally allowed but need usage rules.','Intentional variety can look accidental without explicit hierarchy.','Document page archetypes, use cases and review criteria; enforce in component examples.','Raises premium coherence.','S','P2'),
('R-050','Medium','Design system','No observed visual regression test gate protects the claimed design system.','Small changes can erode hierarchy across 55 pages.','Capture approved states at desktop/tablet/mobile and diff in CI.','Protects brand and accessibility.','M','P2'),
('R-051','Medium','Performance','Several list surfaces fetch full datasets and filter client-side.','Latency and browser load will degrade with real clinic history.','Set page limits, cursors, indexed queries and response budgets.','Supports scale without redesign.','L','P2'),
('R-052','Medium','Performance','Command-center service has a documented N+1 per-clinic query pattern.','Multi-clinic owner experience will slow at precisely the upsell moment.','Replace with aggregate queries/read models and performance tests.','Supports enterprise conversion.','M','P2'),
('R-053','Medium','Reliability','Demo seeding is non-transactional and not concurrency-locked.','Demos can fail visibly or misrepresent product reliability.','Use a durable demo tenant pipeline, lock and readiness check.','Improves sales demo confidence.','M','P2'),
('R-054','Medium','Security','CSP nonce work is deferred.','The browser security posture is weaker than a healthcare product should accept long term.','Complete CSP hardening and include report-only rollout/violation monitoring.','Reduces XSS exposure.','M','P2'),
('R-055','Medium','Security','Trusted proxy handling is configurable only as future debt.','IP-based protections can be bypassed/misapplied in production topology.','Document proxy topology and make trusted hops explicit, tested configuration.','Protects rate limiting and audit accuracy.','S','P2'),
('R-056','Medium','Security','Some public-facing doctor directory reads create PII/exposure risk in the debt register.','Unnecessary data exposure harms patient/staff privacy.','Minimize fields, require intent, audit access and test enumeration resistance.','Reduces privacy risk.','M','P2'),
('R-057','Medium','QA','Tests are heavily service/API oriented; usability acceptance criteria are thin.','A technically correct workflow can still be unusable at the front desk.','Add task-time, error-rate and completion criteria to release evidence.','Improves real-world quality.','M','P2'),
('R-058','Medium','QA','No defect taxonomy tied to clinical severity is evidenced.','Teams may treat patient safety defects as ordinary UI bugs.','Define P0 clinical, P0 privacy, P1 revenue/flow and stop-ship rules.','Improves release decision quality.','S','P2'),
('R-059','Medium','QA','Network/offline UI is claimed but no proof of safe offline mutations is evidenced.','Connectivity messaging without durable sync could create data loss expectation.','Specify offline scope; test queueing/conflicts or remove “sync when we can” claims.','Protects trust in low-connectivity settings.','L','P2'),
('R-060','Medium','Architecture','Capability grants are clinic-agnostic.','Growing groups cannot safely tailor authority by branch.','Introduce clinic-scoped grants before enterprise packaging.','Makes multi-clinic authorization sellable.','L','P2'),
('R-061','Medium','Architecture','Identity unification is deferred while phones are unique across account types.','Staff/patient role overlap creates future support and migration friction.','Design account/profile linkage migration now; avoid new phone-bound constraints.','Preserves platform flexibility.','L','P2'),
('R-062','Medium','Architecture','Duplicate clinic/department prevention is application-level only.','Concurrent writes can create duplicate master data.','Add database uniqueness constraints with data-cleanup migration.','Improves integrity.','M','P2'),
('R-063','Medium','Architecture','The status machine allows scheduled → in consultation directly.','Reception handoff can be bypassed and operational truth becomes ambiguous.','Make bypass an explicit exception reason or remove it for standard flow.','Improves queue integrity.','S','P2'),
('R-064','Medium','Commercial','No quantified ROI calculator or benchmark evidence is visible.','A clinic owner hears aspiration, not a buying case.','Instrument and publish baseline-to-outcome metrics from pilots.','Raises close rate and pricing power.','M','P2'),
('R-065','Medium','Commercial','No clear free trial, assisted pilot or conversion motion is evidenced.','Acquisition funnel and support load cannot be planned.','Test a tightly scoped 30-day assisted pilot with success milestones and conversion trigger.','Creates repeatable GTM learning.','S','P2'),
('R-066','Medium','Commercial','No partner/integration strategy is visible for payments, messaging, labs or accounting export.','Building all adjacent capabilities will dilute the product.','Publish integration principles and choose a small partner roadmap.','Accelerates completeness without ERP drift.','M','P2'),
('R-067','Medium','GTM','The marketing site needs evidence-led claims, security posture and persona-specific proof.','Prospects cannot distinguish promise from implementation.','Create vertical landing pages, workflow demos, proof points, implementation expectations and transparent scope.','Improves qualified demand.','M','P2'),
('R-068','Medium','GTM','No sales enablement for objection handling, comparison or implementation is visible.','Founder-led selling will not scale and claims will drift.','Create battlecards, demo script, security FAQ, migration FAQ and pricing guardrails.','Improves sales consistency.','S','P2'),
('R-069','Medium','Operations','Room/slot capacity and resource scheduling are not operationally mature.','Polyclinic workflows will break even if appointments work.','Keep polyclinics out of target until rooms, resources and capacity policies are validated.','Avoids over-selling.','S','P2'),
('R-070','Medium','Operations','No explicit cancellation/no-show recovery workflow is visible.','Lost capacity is not systematically recovered.','Add outcome reasons, automated rebooking prompt, waitlist offer and owner metric.','Recovers revenue.','M','P2'),
('R-071','Medium','Operations','No explicit queue escalation for clinical urgency is visible.','Priority sorting without policy can be unsafe or arbitrary.','Define limited, auditable urgency flag roles and escalation policy with clinical governance.','Improves safety.','M','P2'),
('R-072','Medium','Documents','Reserved document types are platform-ready but not content-ready.','Users cannot rely on a single document system for routine care.','Prioritize referral, certificate and lab request only after clinical/legal template review.','Improves completeness and stickiness.','M','P2'),
('R-073','Medium','Documents','No document delivery/read receipt workflow exists.','Issued documents may not reach patients or create defensible evidence.','Link document sharing to communications, permissions, expiries and audit receipt.','Improves service and reduces calls.','M','P2'),
('R-074','Medium','Reporting','No data-export/self-service report strategy is evidenced.','Owners and accountants will export manually or abandon the product.','Offer limited role-based CSV/PDF exports and scheduled operational reports.','Improves adoption without becoming ERP.','M','P2'),
('R-075','Medium','Observability','Correlation-ID coverage remains incomplete.','Support investigations will be slow across real user flows.','Finish request correlation and standard operational dashboards/runbooks.','Cuts incident resolution time.','M','P2'),
('R-076','Low','UX','No keyboard shortcut map is surfaced for high-frequency staff actions.','Experienced receptionists cannot reach speed potential.','Add discoverable shortcuts for search, walk-in, check-in, call-in and collect.','Raises throughput.','M','P3'),
('R-077','Low','UX','No saved views or user preferences for queues/worklists are visible.','Users repeatedly reconstruct the same context.','Allow role-safe saved filters and column density preferences.','Improves repeat productivity.','M','P3'),
('R-078','Low','UX','No comparative workload planning view is available to doctors.','Schedule pressure remains reactive.','Add simple day capacity and planned-vs-actual view after scheduling foundations are fixed.','Improves work-life and adoption.','M','P3'),
('R-079','Low','Patient','No patient feedback/review loop follows completed care.','The product misses a low-cost trust and growth signal.','Ask one short, consented experience question after a completed visit.','Creates learning and referral opportunity.','S','P3'),
('R-080','Low','Patient','Patient education content is absent.','A portal can feel transactional rather than caring.','Support clinic-authored pre/post-visit instructions only, not medical advice generation.','Raises perceived care quality.','M','P3'),
('R-081','Low','Design','Iconography is sourced from a standard set but a documented usage guide is not visible.','Semantic variation can become inconsistent.','Publish icon meaning, status pairing and do-not-use rules.','Improves learnability.','S','P3'),
('R-082','Low','Design','Mobile responsive behavior is asserted but not independently evidenced by visual testing.','Desktop-first layouts may fail in patient contexts.','Create breakpoint acceptance snapshots and real-device pilot checks.','Reduces mobile churn.','M','P3'),
('R-083','Low','Engineering','External font fetching makes builds environment-sensitive.','Builds can fail in restricted networks, as this audit observed.','Self-host approved fonts or provide robust local fallback in the build pipeline.','Improves reproducibility.','S','P3'),
('R-084','Low','Engineering','The repository README remains the default Next.js starter text.','New engineers and evaluators lack an accurate operating entry point.','Replace with architecture, local setup, quality gates, test data and production runbook pointers.','Reduces onboarding cost.','S','P3'),
('R-085','Low','Engineering','Production readiness claims are dispersed across many documents.','Teams can miss the actual release truth.','Create one versioned release dashboard with evidence links and owner/date for each gate.','Improves governance.','S','P3'),
('R-086','Low','Product','Doctor profile contains several “coming soon” tiles.','It dilutes professional credibility without advancing care.','Keep profile focused on bookable, verified attributes; defer aspirational features from live UI.','Improves trust.','S','P3'),
('R-087','Low','Product','Patient insurance settings expose a non-functional pathway.','Creates confusion in a market where coverage expectations vary.','Remove until supported, or clearly explain the no-insurance product policy.','Cuts support requests.','S','P3'),
('R-088','Low','Product','The public directory lacks specialty/location/availability filters.','Discovery does not serve patient intent at scale.','Add search ranking and filters once the target geography and provider supply model are defined.','Improves booking conversion.','M','P3'),
('R-089','Low','Product','No structured review-management workflow is visible despite doctor ratings.','Ratings can appear stale or unmanaged.','Define source, moderation, response and consent policy before using ratings commercially.','Protects brand trust.','M','P3'),
('R-090','Low','Product','Document categories are broad without a patient-friendly naming strategy.','Patients may not recognize clinical artifacts.','Use clear consumer labels, summaries and progressive disclosure in portal.','Improves comprehension.','S','P3'),
('R-091','Low','Product','No ROI nudges at clinic milestones are visible.','Owners may not perceive the value created by operations improvement.','Show monthly avoided no-shows, recovered balances and rebooking outcomes when data is reliable.','Supports retention.','M','P3'),
('R-092','Low','Product','No in-product learning/help pattern is visible for first-time workflows.','“Zero training” is an unproven aspiration.','Use contextual, dismissible guidance and role-based first-run tasks.','Reduces onboarding friction.','M','P3'),
('R-093','Low','QA','No production synthetic monitoring of public booking/OTP is evidenced.','Failures may be discovered by patients first.','Run consent-safe synthetic checks and alert on booking/OTP degradation.','Improves reliability.','M','P3'),
('R-094','Low','QA','No load test thresholds for morning queue spikes are documented.','Performance risk is unmeasured at the highest-stress clinical moment.','Benchmark queue/search/checkout under representative concurrent load.','Protects clinic flow.','M','P3'),
('R-095','Low','QA','No accessibility statement or public conformance target is visible.','Enterprise buyers cannot assess inclusion maturity.','Publish supported standard, known limitations and feedback route after audit.','Improves procurement readiness.','S','P3'),
('R-096','Low','Architecture','Event platform has no customer-visible consumer product yet.','Engineering sophistication has little commercial narrative.','Use events first for auditable communications and operational alerts, not generic platform messaging.','Converts foundation into value.','M','P3'),
('R-097','Low','Architecture','There is no explicit data archival strategy visible for high-volume clinical history.','Long-term performance and retention cost are unknown.','Define retention/archival policies, query boundaries and legal review.','Supports scale responsibly.','M','P3'),
('R-098','Low','Commercial','No formal enterprise SLA/support package is ready.','Enterprise tier risks being only a feature label.','Do not sell Enterprise until security review, support SLAs, governance and multi-clinic depth exist.','Avoids churn and reputational damage.','S','P3'),
('R-099','Low','GTM','No customer reference program is evident.','Auriva lacks proof beyond product narrative.','Recruit 3–5 design partners with explicit outcome measurement and reference rights.','Builds believable demand.','M','P3'),
('R-100','Low','Strategy','The roadmap sequence is sound in intent but lacks decision gates tied to customer evidence.','Teams can build elegant features without commercial pull.','Require a validated problem, target segment, success metric and stop condition before each milestone.','Improves capital allocation.','S','P3'),
]
recs.extend(more)
assert len(recs) == 100

screens = [
('Marketing home','/','6.0','Clear category story and visual system.','“Healthcare OS” promise exceeds present product breadth.','Use a concrete front-desk-to-cash proof story, pilot outcomes and scope boundaries.'),
('About','/about','5.5','Brand mission can build empathy.','Mission is not buyer evidence.','Add team credibility, clinical governance and customer proof.'),
('Platform','/platform','5.0','Pillar framing is coherent.','Claims need implementation-level substantiation.','Map each promise to shipped capability or remove it.'),
('Solutions','/solutions','4.8','Possible segment orientation.','Broad segments hide poor beachhead focus.','Create one vertical page per validated segment only.'),
('Industries','/industries','4.5','Signals ambition.','Cross-specialty claims exceed current specialized workflow depth.','Limit to validated specialties.'),
('Pricing','/pricing','3.0','A pricing surface exists.','Commercial package is not finalized.','Do not expose until plans, limits and support are sellable.'),
('Security','/security','5.0','Security concern is acknowledged.','Claims need a procurement-ready evidence pack.','Publish controls, scope and independent review status.'),
('Trust','/trust','5.0','Trust is a useful brand theme.','Trust cannot substitute for communications, reliability and governance.','Use audited operational proof.'),
('Compliance','/compliance','4.2','Signals seriousness.','Compliance posture is not evidenced as complete.','Obtain legal review and state exact jurisdictional position.'),
('Customers','/customers','3.5','Correct destination for proof.','No demonstrated reference base is evident.','Populate only with attributable pilot outcomes.'),
('Book demo','/book-demo','5.5','Appropriate enterprise CTA.','No qualification/expectation framework is shown.','Add target-fit questions and implementation promise.'),
('Contact sales','/contact-sales','5.0','Clear route to human help.','Sales cannot answer pricing/migration confidence yet.','Equip sales before demand generation.'),
('Get started','/get-started','4.5','Supports self-serve aspiration.','First-value and setup readiness are unproven.','Gate through assisted pilot until activation is proven.'),
('Public booking','/book/[doctorId]','5.8','Real slot selection is a good foundation.','No intake, consent, payment or reminder loop.','Complete booking-to-arrival lifecycle.'),
('Start','/start','5.0','Account entry is explicit.','Patient activation depends on unshipped production OTP.','Block external launch until delivery is live.'),
('Login','/login','5.5','Role-aware entry and security states appear considered.','Lint errors sit in this high-risk screen; selection adds cognitive load.','Clean quality gate and reduce first-time ambiguity.'),
('Register organization','/register-org','5.0','Organization model is real.','No clear guided operational setup/migration.','Add stepped setup with readiness confirmation.'),
('Workspace selector','/workspace','6.0','Multi-workspace model is a real differentiator.','Value is niche before customers are multi-clinic.','Make solo path exceptionally simple; reveal complexity progressively.'),
('Admin home','/admin','5.3','Owner/admin surface exists.','Admin capability outpaces commercial readiness.','Focus on actionable clinic health, not platform furniture.'),
('Command center','/admin/command-center','5.2','Shows operational attention.','Analytics and resolution actions are thin.','Prioritize outcomes, owner and one-click recovery.'),
('Departments','/admin/departments','5.0','Supports organization structure.','Not a daily buyer value driver.','Keep out of launch demo unless polyclinic fit is proven.'),
('Admin events','/admin/events','4.5','Good engineering observability.','Customer-facing admin value is unclear.','Keep internal unless audit/replay is a paid need.'),
('Admin releases','/admin/releases','3.5','Useful internal release construct.','Does not belong in normal clinic operating narrative.','Protect behind platform-admin role.'),
('Admin settings','/admin/settings','4.8','Basic organization management is necessary.','Plan/clinic nav remains “Soon.”','Remove incomplete navigation or finish it.'),
('Admin setup','/admin/setup','5.2','Checklist is directionally correct.','No evidence of activation completion/value metric.','Instrument and make blockers actionable.'),
('Clinic home','/clinic','5.8','Useful consolidated solo surface.','May become a crowded hybrid as features grow.','Define strict solo cockpit information architecture.'),
('Clinic calendar','clinic dashboard','5.0','Calendar foundation exists.','Scheduling depth lacks recurrence/resources.','Do not promise polyclinic operations yet.'),
('Clinic consultation','clinic dashboard','5.5','Shared consult components reduce drift.','Structured clinical model remains immature.','Advance encounter model before expanding templates.'),
('Clinic plan','clinic dashboard','4.0','Plan intent aligns with strategy.','Commercial/treatment planning maturity is incomplete.','Deliver the full plan loop before promotion.'),
('Clinic team','clinic dashboard','5.5','Membership/RBAC foundation is sound.','Nurse/technician daily work is missing.','Align role packaging to real actions.'),
('Staff home','/staff','5.5','Reception surface addresses operational flow.','Navigation can fragment queue/calendar/billing context.','Create a unified front-desk day workspace.'),
('Staff dashboard','/staff/dashboard','3.5','Contains some useful search affordance.','Documented as orphaned/unlinked.','Remove or make it a deliberate landing page.'),
('Queue board','/staff/queue','6.5','Strong stage ownership, wait aging, one-action-per-stage.','No reassign, capacity recovery, global patient search or full calendar context.','Keep as flagship; complete its exception paths.'),
('Walk-in','/staff/walkin','6.0','Minimal capture and identity resolution are valuable.','Family disambiguation and identifier use are weak.','Use Health ID/DOB and explicit duplicate resolution.'),
('Billing desk','/staff/billing','5.7','Checkout, partial/split payment and corrections show real substance.','No online payment, weak search, tax inactive.','Finish patient payment and compliant billing loop.'),
('Staff calendar','/staff/calendar','4.8','Gives a day view.','Day-only, fixed hours, no scheduling power.','Develop before positioning as scheduling solution.'),
('Lab worklist','/staff/lab','4.8','Basic order/result flow exists.','Search, clinical review and integration depth are weak.','Keep narrow or invest in trusted lab workflow.'),
('Doctor Today','/doctor','6.3','“Who is next?” is a strong mental model.','Pause booking is non-functional; no deeper exceptions.','Retain focus and remove false controls.'),
('Doctor Workbench','/doctor/workbench','6.4','Three-column context/queue/consult concept is compelling.','Free-text data, autosave confidence and placeholder context limit clinical trust.','Make this the validated clinical flagship.'),
('Doctor Patients','/doctor/patients','4.6','Derived patient list is a start.','Name-only client search and no risk/follow-up facets.','Add secure server search and useful clinical filters.'),
('Doctor schedule','/doctor/schedule','5.0','Day/week/month visibility and blocks exist.','Requests and mature scheduling are absent.','Make scheduling reliable before adding analytics.'),
('Doctor practice','/doctor/practice','3.8','Profile intent is clear.','Many “coming soon” panels make it feel unfinished.','Strip to real booking/profile management.'),
('Doctor profile','/doctor/profile','5.0','Credentials and availability matter.','Awards, memberships and signature are incomplete tiles.','Only show verified, usable attributes.'),
('Patient home','/patient','6.4','Warm, mobile-first framing and next-visit summary are sound.','It cannot proactively communicate or resolve payment.','Turn it into an active care companion.'),
('Patient booking','/patient/book','5.2','Simple discovery and booking path.','Client-side directory search lacks filters and scale.','Build intent-led specialty/location/availability discovery.'),
('Patient doctor profile','/patient/doctors/[id]','5.6','Supports profile-to-slot conversion.','No intake, consent, deposit or expectation setting.','Use pre-visit preparation as conversion advantage.'),
('Patient care','/patient/care','4.5','Care concept aligns with product vision.','The longitudinal care proposition is still thin.','Anchor it in real plans, Rx and communication.'),
('Patient records','/patient/records','5.2','Timeline/Rx/Bills/Tests grouping is understandable.','No record search; depth depends on unstructured visit data.','Ship timeline and metadata search together.'),
('Patient family','/patient/family','5.5','Family profile switching suits India-first use.','Consent/proxy controls are immature.','Add clear guardian and sharing governance.'),
('Patient profile','/patient/profile','5.3','Health summary editing is useful.','Critical safety data quality/review status is unclear.','Add source and last-reviewed signals.'),
('Patient settings','/patient/settings','4.2','Session/privacy intent is appropriate.','Multiple settings remain “coming soon.”','Keep only controls that work.'),
('Patient You','/patient/you','5.0','Clear account framing.','Insurance is shown without supporting workflow.','Remove unimplemented financial claims.'),
('Print document','/print/document/[id]','5.5','Central document-rendering foundation is valuable.','Patient delivery and document breadth are incomplete.','Complete delivery/audit before marketing documents.'),
('Print prescription','/print/prescription/[id]','5.0','Print supports a real clinic behavior.','Rx safety and structured content still lack maturity.','Prioritize signing/amendment and safeguards.'),
('Invite join','/join/[token]','5.4','Provisioned-team onboarding is real.','Role productivity is uneven after acceptance.','Link invitation to role-specific first-day checklist.'),
('Change password','/change-password','5.4','Forced password change is a strong safety control.','No self-service recovery reduces support resilience.','Provide secure reset before scale.'),
]

def add_title(d):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(96); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('INDEPENDENT INSTITUTIONAL'); font(r,11,True,PINE)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8)
    r=p.add_run('Product Readiness Assessment'); font(r,30,True,INK)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18)
    r=p.add_run('Auriva Healthcare Operating System'); font(r,17,False,PINE)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('Commercial-launch audit | 23 July 2026 | Evidence reviewed: code, tests, product artefacts, release documents, and current competitor materials'); font(r,9,False,MUTED)
    callout(d, 'MANDATE', 'This is not an implementation sign-off. It is an external product, clinical-workflow, UX, business and engineering review. The assessment challenges internal readiness claims and treats missing evidence as a risk.', 'FBE5E3')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(120)
    r=p.add_run('FINAL LAUNCH DECISION: ⚠ NEEDS SIGNIFICANT WORK'); font(r,15,True,RED)
    pagebreak(d)

def add_method(d):
    h(d, 'Audit scope, method and limits', 1)
    para(d, 'This assessment is evidence-led. It reviewed 55 application pages, 110 API route handlers, 74 automated test files, the Prisma domain model, operating/debt/release documentation, and five internal Product Office workflow audits. A fresh test run completed successfully; TypeScript type-check completed successfully. The project lint command failed with 43 errors and 51 warnings. The production build could not complete in this restricted environment because Next font fetching failed; this is classified as a build reproducibility risk, not proof the product cannot build in a connected production environment.')
    para(d, 'The audit attempted to render the local application in the controlled browser. The browser could not reach the local development server, therefore visual findings are source- and component-led rather than a pixel-level rendered inspection. Scores are intentionally conservative where runtime evidence was unavailable.')
    h(d, 'Evidence hierarchy', 2)
    for x in ['Runtime/automated check result', 'Implemented code/schema/API route', 'Product and release artefact', 'Official competitor material', 'Inference or strategic recommendation']:
        bullet(d, x)
    h(d, 'Evidence sources', 2)
    for x in ['Internal: docs/AURIVA-PRODUCT-BIBLE.md; docs/AURIVA-RELEASE-1X-ROADMAP.md; docs/RELEASE-CANDIDATE.md; docs/technical-debt.md; docs/product-office-audit/*; docs/knowledge-base/*; src/app, src/components, src/services and prisma/schema.prisma.', 'External official competitor sources current at audit date: Jane feature/booking pages; Cliniko online payments; SimplePractice client portal and billing automation; Halaxy calendar; Practo Ray; DocPulse; Clinicea features; Healthie pricing. URLs are listed in Appendix A.']:
        bullet(d, x)

def add_exec(d):
    h(d, '1. Executive Summary', 1)
    callout(d, 'BOARD CONCLUSION', 'Auriva is not launch-ready as a commercial Healthcare Operating System. It is a serious, well-architected, early product with an unusually coherent encounter-to-cash thesis; it is not yet a dependable, purchasable clinic operating system across the patient, revenue, clinical and launch-support lifecycle.', 'FBE5E3')
    para(d, 'The most valuable thing Auriva has is not its breadth. It is the decision to make a clinic day move through one shared appointment state machine: booking/walk-in → waiting → consultation → checkout. That is a credible operational core. The problem is that almost every purchase-critical loop around it is incomplete: patient activation depends on unproven production OTP; the product does not communicate outside the app; patients cannot pay online; intake and consent are absent; scheduling is shallow; analytics are thin; clinical data is largely unstructured; pricing is not finalized; and release gates remain open.')
    table(d, ['Dimension','Score','Status','Board finding'], scores, [2550,700,900,5210], 8.5)
    h(d, 'Two-page executive conclusion', 2)
    para(d, 'The internal materials repeatedly use “mature,” “production-ready” and “release candidate.” Those terms are defensible only for a constrained single-clinic, staff-assisted encounter-to-cash pilot after infrastructure gates are closed. They are not defensible for a broad commercial launch, a patient-led product, a multi-clinic platform or an enterprise edition. The report therefore rejects a launch narrative that outruns the finished experience.')
    para(d, 'Auriva should not respond by adding every competitor feature. That would violate its own guardrails and lose its best strategic asset: operational simplicity. It should instead complete the smallest number of end-to-end loops that make a selected clinic buyer visibly better off every day: reliable identity/communications, pre-visit readiness, front-desk exception management, structured clinically safe documentation, checkout/payment, follow-up/next booking, and owner proof of value.')
    h(d, 'The five non-negotiable launch blockers', 2)
    for x in ['Production OTP delivery, delivery observability and patient-login proof.', 'Backups, restore drill, environment validation, monitoring/alerting and operational ownership.', 'A clean quality gate: lint, type check, tests and production build must be reproducible.', 'Consent-based outbound communication and at least one patient-reachable reminder/document pathway.', 'A sellable commercial package: chosen beachhead, finalized price/limits, assisted onboarding and pilot success definition.']:
        bullet(d,x)
    h(d, 'Final verdict', 2)
    para(d, '⚠ Needs Significant Work. Auriva is appropriate for tightly controlled design-partner pilots only after P0 work is closed. It should not be called “Launch Ready,” should not sell Enterprise, and should not promise a complete Healthcare OS in market messaging.')

def add_vision_market(d):
    h(d, '2. Product Vision Review', 1)
    para(d, 'The vision is unusually clear in prose: an India-first independent clinic should run a calm room, capture cash reliably and make patients feel cared for. The six-pillar boundary is disciplined, and the explicit refusal to become HRMS/ERP is strategically correct. The positioning, however, is too large for the delivered product. “Healthcare Operating System” invites comparison with products that already offer mature communications, online payment, forms/consent, reports, recurring scheduling, telehealth and broader clinical record functionality.')
    table(d,['Question','Board answer','Required correction'],[
        ('Is the vision clear?','Yes. The room-flow and revenue-leak thesis is memorable.','Preserve it; turn it into an observable promise.'),
        ('Is positioning obvious?','Partly. The category is clear, the buyer and beachhead are not.','Name a clinic type, specialty and daily outcome.'),
        ('Would customers understand it?','They understand appointments/queue/billing; “OS” can confuse.','Lead with a concrete before/after workflow.'),
        ('Would investors understand it?','They will see ambition but ask why incumbents do not win.','Show segment evidence, retention mechanism and unit economics.'),
        ('Would doctors understand it?','They will understand the workbench but reject empty promises.','Demonstrate safety, speed and post-visit continuity.')], [1900,3100,4360],8.5)
    h(d, '3. Target Customer Review', 1)
    table(d,['Segment','Fit now','Board judgment','Why'],[
        ('Solo GP/consulting doctor','Amber','Pilot only','Core appointment/consult/billing works, but patient engagement and discovery are thin.'),
        ('Physiotherapy','Amber → Green after C1/C2','Best beachhead candidate','Multi-session plans and follow-up leakage map directly to roadmap.'),
        ('Dentistry','Amber → Green after C1/C2','Strong but not now','Treatment plans, procedures, consent and imaging/doc needs are not ready.'),
        ('Psychology','Red','Avoid initially','Needs intake, consent, secure messaging, telehealth and specialty documentation.'),
        ('Nutritionists','Red','Avoid initially','Healthie/Jane strength is care plans, messaging and remote engagement; Auriva is not there.'),
        ('Single-site multi-doctor clinic','Amber','Controlled pilot','Best fit for queue/reception thesis once P0s close.'),
        ('Polyclinic/chain','Red','Do not sell','Rooms/resources, deep multi-clinic, enterprise controls and reporting are not mature.')], [1900,1000,1800,4660],8.2)
    h(d, '4. Market Readiness', 1)
    para(d, 'Auriva’s potential differentiation is operationally coherent, not feature breadth: a shared appointment state machine, India-first phone/family/UPI assumptions and a deliberately warm patient experience. It is not differentiated enough today to overcome competitor table stakes. Jane connects online booking, reminders, intake and payments; Cliniko can collect online payment and issue payment links; SimplePractice automates billing document delivery and portal payments; Halaxy connects online bookings, reminders, records, invoicing and reporting; Practo Ray promotes SMS reminders, online payment and treatment plans; Clinicea advertises automated consent, reminders, packages, telehealth and patient access. Auriva must win a narrower workflow, not claim parity.')
    table(d,['Competitor','What it establishes as table stakes','Auriva gap / response'],[
        ('Cliniko','Mobile-friendly online payments, payment links and automated invoices.','No online payment; prioritize India-appropriate UPI/payment flow.'),
        ('Jane','Online booking, reminders, waitlist, intake/consent, payments, secure portal and scheduling depth.','Do not compete broadly; match pre-visit + communication basics for chosen segment.'),
        ('Healthie','Telehealth, scheduling, business tools and wellness/team workflows.','Avoid remote-care/nutrition market until care plans and messaging mature.'),
        ('Halaxy','Online booking, email/SMS reminders, patient records, invoices and financial reports.','Communication and owner reporting are non-negotiable gaps.'),
        ('SimplePractice','Portal, online payment, billing automation, documentation/search/filter depth.','Patient self-service and information retrieval are thin.'),
        ('Practo Ray','India-specific reminders, online payment, EMR/treatment plans and mobile access.','Auriva must prove lower friction/front-desk superiority, not generic parity.'),
        ('DocPulse','Broad integrated HMS/clinic suite incl. EMR, billing, labs, pharmacy, reminders, analytics.','Do not chase hospital breadth; own independent clinic flow.'),
        ('Clinicea','Consent, plans/packages, telehealth, communication, multi-location and deep specialty tooling.','Auriva needs a beachhead and simpler proof, not a feature race.')], [1250,4200,3910],7.7)

def add_journeys(d):
    h(d, '5. Complete User Journey Audit', 1)
    journey = [
        ('Account/clinic creation','Organization/workspace foundation is real.','No proven first-value, import or guided operational configuration.','Instrument setup and offer assisted pilot onboarding.'),
        ('Login/identity','Staff sessions, forced password change and OTP design are thoughtful.','Production OTP delivery is not complete; patient beta is blocked.','Close R-001 before exposing patient journey.'),
        ('Booking','Slot booking and server conflict checks exist.','No forms, consent, deposit, reminder or waitlist recovery.','Complete pre-visit loop.'),
        ('Walk-in/reception','Fast minimal capture and status machine are strong.','Search/disambiguation, reassign and exception recovery are weak.','Surface Health ID/DOB; add transfer.'),
        ('Queue/consult','Clear state ownership minimizes ambiguity.','Direct consult bypass and shallow escalation can create operational drift.','Guard exceptions and display recovery.'),
        ('Clinical documentation','Shared workbench and factual safety chips are promising.','Free-text encounter data, no signed amendment model or clinical governance.','Structure core data and govern releases.'),
        ('Prescription/labs','Print and basic order/result paths exist.','No full medication safety, clinical review or integration depth.','Avoid over-claiming clinical maturity.'),
        ('Checkout/payment','Invoice/payment/correction architecture is a real strength.','No online payment, tax active fields or unified service-to-charge proof.','Close payment and service capture loops.'),
        ('Records/documents','Timeline/document platform foundation is good.','Search, delivery, patient-friendly access and longitudinal completeness are weak.','Ship timeline + communication together.'),
        ('Notifications/follow-up','In-app notification exists.','No reachable notifications; no automated follow-up.','This is a launch blocker, not polish.'),
        ('Settings/logout','Role/session handling is robust in intent.','Several surfaces expose “Soon,” confusing real vs planned actions.','Hide/defer incomplete controls.')]
    table(d,['Stage','What works','Break/failure point','Recommendation'],journey,[1450,2650,2700,2560],7.6)
    h(d, '6. Role-Based Journey Review', 1)
    table(d,['Role','Ease / cognitive load','Breakpoints','Board recommendation'],[
        ('Doctor','Good queue-to-consult focus; workbench can be compelling.','Search weak; clinical structure/safety immature; placeholders harm trust.','Make Workbench one validated, high-confidence clinical surface.'),
        ('Receptionist','Best-designed role conceptually; stage ownership is intuitive.','Search, reassignment, calendar context, communications and payment recovery are incomplete.','Make reception the beachhead and test it under real traffic.'),
        ('Clinic owner','Has command-center/team/billing foundation.','Cannot yet see a measurable business case or mature reports.','Sell outcomes dashboard, not administration.'),
        ('Patient','Warm framing and family model are promising.','Login, communication, payments and after-visit support are weak.','Build active, not passive, patient relationship.'),
        ('Administrator','Strong RBAC/audit/event scaffolding.','Some admin UI is internal/platform-like; enterprise controls incomplete.','Keep admin scope narrow until enterprise needs are real.')], [1250,2500,2800,2810],8)
    h(d, '7. UI Review', 1)
    para(d, 'The implemented design system is more mature than the business product. It has semantic colors, light/dark tokens, common components, focus-visible rules and reduced-motion handling. The intended Pine/Honey separation can be a distinct brand asset. The weakness is product truth: polished cards and intentional copy cannot compensate for disabled “Soon” navigation, informational-only controls, and operational journeys that stop before patient communication/payment. A premium product feels complete in the moments people pay for, not merely consistent in its component code.')
    h(d, '8. UX Review', 1)
    para(d, 'Mental model strength: the appointment state machine gives each role a visible handoff. This is excellent and should remain sacrosanct. Mental model failure: user intent is repeatedly acknowledged but not completed—pause booking, insurance, upcoming capabilities, communications, advanced schedule management. Each is a small expectation breach. The UX roadmap should prioritize recovery and closure over decorative refinement.')
    h(d, '9. Healthcare Workflow Review', 1)
    para(d, 'A real receptionist can plausibly process a walk-in, move a patient through the queue and collect payment. That is a meaningful vertical slice. A real clinic, however, also needs certainty before arrival, handling for late/absent/changed doctors, room/capacity decisions, patient reminders, rebooking, consent, follow-up and owner visibility. The product handles the center of a clinic day better than its edges. Commercially, the edges are where no-shows, liability and churn happen.')

def add_business_quality(d):
    h(d, '10. Business Model Review', 1)
    para(d, 'There is an edition shape (Solo / Professional / Enterprise) but no approved pricing. That is not pricing readiness. Feature gates alone are not packaging. A viable near-term commercial offer is: an assisted, single-clinic Professional pilot for one validated specialty, with a fixed implementation fee, named service limits, core staff seats, committed support and a 90-day outcome review. Do not offer Enterprise until shared rate limits, clinic-scoped grants, multi-clinic analytics, security/compliance evidence, support SLAs and proven customer operations exist.')
    table(d,['Revenue lever','Readiness','Board view'],[
        ('Subscription tiers','Red','Shape exists, commercial terms do not.'),('Implementation/onboarding','Red','No mature migration or success operating model.'),('Communications usage','Red','High potential after consented delivery exists.'),('Payment processing','Red','High potential but no patient payment capability.'),('Treatment plans/packages','Amber','Strong future revenue lever; needs full workflow first.'),('Multi-clinic/enterprise','Red','Do not monetize prematurely.'),('Analytics/reporting','Amber','Renewal lever once decision-grade metrics exist.')],[2500,1000,5860],8.5)
    h(d, '11. Quality Assurance Audit', 1)
    para(d, 'Positive evidence: 74 test files, 704 test declarations, 110 API routes, centralized status transitions, audit/event seams, health/readiness endpoints and documented operational debt. Negative evidence matters more for launch: the current lint run fails with 43 errors/51 warnings; UI-level E2E proof is not evidenced; production build reproducibility is sensitive to remote font fetching in this environment; single-instance rate limiting/reminder sweeps and unpaged list APIs are documented limitations. The test count is not a release decision. Gate on scenario coverage, severity and clean commands.')
    h(d, '12. Accessibility Audit', 1)
    para(d, 'Positive signals include focus-visible styling, several ARIA labels, semantic status labels, reduced-motion CSS, shared controls and dark-mode tokens. Gaps: no automated accessibility audit evidence, no keyboard workflow certification, no screen-reader transcript evidence, no touch-target audit and no contrast evidence across all dynamic states. Accessibility is promising in implementation intent, not verified in launch assurance.')
    h(d, '13. Design System Audit', 1)
    para(d, 'The system has genuine reuse: shared controls, semantic CSS variables, dark mode, state components and a clear staff/patient tone. This is a comparative strength. It needs a governed token inventory, visual regression tests, icon/status usage rules and fewer product-level exceptions. The design system should protect the operational core, not create a false sense of maturity around unfinished workflows.')
    h(d, '14. Product Maturity', 1)
    table(d,['Classification','Board assessment'],[
        ('MVP','Yes: a real staff-assisted encounter-to-cash slice exists.'),
        ('Friendly-clinic beta','Conditional: only after P0 launch/infrastructure, quality and communication gates close.'),
        ('Launch candidate','No: commercial, patient and operational completeness is insufficient.'),
        ('Production ready','Only narrow internal/single-clinic components, not the commercial product.'),
        ('Enterprise ready','No: do not market or sell it.')],[2400,6960],9)

def add_roadmap_launch(d):
    h(d, '15. Future Roadmap Review', 1)
    para(d, 'The stated sequence—Treatment Planning → Procedure Management → Clinical Timeline → Prescription Platform → Communication—is strategically sound but incomplete as a launch plan. Communication must move ahead of broad patient/pilot claims because it is the delivery mechanism for the existing promise. Treatment planning is the best commercial growth step only if the beachhead is physio/dental; otherwise it is premature specialization. Do not add generic “revenue platform” breadth. Make services-to-invoice, payment, rebooking and outcome reporting reliable first.')
    table(d,['Roadmap item','Board verdict','Required change'],[
        ('Revenue platform','Keep narrow','Complete payment, tax decision, service capture, correction and owner reporting—not accounting.'),
        ('Treatment planning','Prioritize for selected specialty','Require plan-to-session-to-payment-to-next-booking end-to-end metric.'),
        ('Procedure management','Do after plans','Avoid specialty hardcoding; validate session model with design partners.'),
        ('Clinical timeline','Accelerate','Make it doctor pre-consult and patient record truth, with performance/search.'),
        ('Prescription platform','Accelerate with governance','Structured Rx, signing/amendment and safety checks—not AI diagnosis.'),
        ('Communication platform','Move to launch-critical','Consent, templates, delivery, preference and audit is required before public beta.')],[1900,2300,5160],8)
    h(d, '16. Launch Readiness', 1)
    callout(d,'LAUNCH BLOCKED','The product should not launch broadly. The release documentation itself lists unchecked production database, backup/restore, production OTP/SMS, production configuration and monitoring/alerting gates. This audit adds a failing lint gate, absent commercial packaging, no outbound patient communication, no online patient payment and insufficient E2E/clinical governance evidence.','FBE5E3')
    h(d, '17. Commercial Readiness', 1)
    para(d, 'Would clinics pay? A narrow group might pay for an assisted pilot if Auriva proves fewer missed handoffs, faster walk-ins and fewer uncollected visits. Would they switch? Not yet at scale: data migration, communication, payment and recurring scheduling are table stakes. Would users recommend it? Receptionists could recommend the queue board; patients will not recommend a portal that does not reliably reach them, accept payment or guide them after the visit. Stickiness will come from longitudinal care plans, documents/communication, a trusted patient record and owner ROI—not from the admin architecture.')
    h(d, '18. SWOT Analysis', 1)
    table(d,['Strengths','Weaknesses','Opportunities','Threats'],[
        ('One status machine across roles; clear product guardrails; strong billing/audit foundation; warm patient design intent.','Unfinished patient communication/payment/pre-visit loop; weak commercial readiness; shallow scheduling/search/analytics; clinical data debt.','Own India-first single-site clinic flow; win multi-visit specialty; prove outcomes; turn communications/timeline into retention.','Feature-complete incumbents; scope dilution into ERP/HIS; clinical/privacy incident; no reference customers; expectation debt.')],[2340,2340,2340,2340],7.5)
    h(d, '19. Risk Register', 1)
    risks=[('Business','Red','No beachhead/pricing/ROI','Cannot sell repeatably','CPO','R-008/R-028/R-064'),('Product','Red','Patient loop stops at portal','Low adoption/no-show value unrealized','VP Product','R-005/R-009/R-043'),('UX','Amber','Front-desk exception paths missing','Stress and workaround behavior','UX Director','R-013/R-019/R-039'),('Technical','Red','Quality gates fail; single-instance dependencies','Regression/reliability risk','CTO','R-003/R-006/R-007'),('Clinical','Red','Unstructured records, no governance/consent','Safety and liability risk','Clinical Safety Board','R-004/R-015/R-017'),('Commercial','Red','Enterprise claims ahead of capability','Churn/reputation risk','GTM Lead','R-030/R-098'),('Security/privacy','Amber','Compliance pack and scale hardening incomplete','Procurement/security risk','Security Lead','R-020/R-054/R-056')]
    table(d,['Risk class','Status','Risk','Consequence','Owner','Mitigation'],risks,[1000,650,2450,2100,1350,1810],7.7)

def add_recs(d):
    h(d, '20. Top 100 Recommendations', 1)
    para(d, 'The register is sequenced by launch impact. “Critical” means a launch/pilot gate; “High” means the first commercial roadmap; “Medium” means required to protect adoption and scale; “Low” means valuable only after the core is proven. Complexity: S ≤1 sprint, M 1–2, L 2–4, XL multi-milestone.')
    groups=[('Critical','P0'),('High','P1'),('Medium','P2'),('Low','P3')]
    for pri, code in groups:
        h(d, f'{pri} — {code}', 2)
        rows=[]
        for r in recs:
            if r[1]==pri:
                rows.append((r[0],r[2],r[3],r[4],r[5],r[6],r[7]+' / '+r[8]))
        batches=[rows[i:i+5] for i in range(0,len(rows),5)]
        for n,batch in enumerate(batches,1):
            if len(batches)>1:
                p=d.add_paragraph(); p.paragraph_format.space_after=Pt(3); r=p.add_run(f'{pri} recommendation detail {n} of {len(batches)}'); font(r,8.5,True,MUTED)
            table(d,['ID','Area','Problem','Impact','Recommendation','Business value','Complexity / priority'],batch,[620,850,2020,1600,2090,1480,700],7.6)
            if n < len(batches): pagebreak(d)

def add_screens(d):
    h(d, '21. Screen Review', 1)
    para(d, 'Scores are source-led because the controlled browser could not reach the local development server. They assess implemented purpose, connected capability, known limitations and visible route/component intent—not a pixel-perfect live render.')
    chunks=[screens[i:i+10] for i in range(0,len(screens),10)]
    for idx,ch in enumerate(chunks,1):
        h(d, f'Screen inventory {idx} of {len(chunks)}', 2)
        table(d,['Screen','Route','Score','Strength','Weakness','Missing information / suggestion'],ch,[1320,1400,520,2000,2180,1940],6.8)

def add_emotion_delight(d):
    h(d, '22. User Emotion Journey', 1)
    table(d,['Moment','Likely emotion now','Why','Design response'],[
        ('Discover / buy','Curious then skeptical','A strong story meets a broad, unfinished product claim.','Make promise specific and evidence-led.'),
        ('First setup','Hopeful then uncertain','Organization capability exists; activation/migration path is not proven.','Guide to a first successful clinic day.'),
        ('Reception rush','Potentially confident','Board makes next state visible.','Add search, reassign, escalation and calendar context.'),
        ('Doctor consult','Focused then cautious','Workbench is structured visually, but record/safety depth is limited.','Make saving, signing and context trustworthy.'),
        ('Checkout','Efficient in person','Billing architecture is solid.','Enable patient payment and clear service linkage.'),
        ('At home patient','Neglected','No external reminder/document/payment/care loop reaches them.','Create a consented after-visit relationship.'),
        ('Owner review','Underwhelmed','Activity is visible; outcomes/ROI are not.','Show how Auriva improved the month.')],[1400,1750,3200,3010],7.8)
    h(d, '23. Customer Delight Opportunities', 1)
    for x in ['A “clinic day closed” moment: every completed encounter shows collected/owed, unresolved exceptions and tomorrow’s risk—not generic dashboard cards.', 'One-tap rebook at checkout tied to a plan/follow-up, with patient confirmation sent automatically.', 'Patient after-visit card: next action, document, payment status, contact path and reassuring language.', 'Reception recovery panel: “Dr. late, 4 waiting >30m, 2 alternative slots” with controlled actions, not an alarm wall.', 'Owner value note: “This month Auriva helped recover X follow-ups / Y outstanding balances”—only when data is reliable.', 'Structured prescription that prints beautifully, is clearly signed/versioned and lands safely in the patient record.']:
        bullet(d,x)
    h(d, '24. Final Product Report Card', 1)
    report=[('Product','5.2','Red'),('Design','6.8','Amber'),('UX','5.4','Amber'),('Engineering','5.8','Amber'),('Architecture','6.5','Amber'),('Scalability','4.2','Red'),('Business','3.7','Red'),('Commercial','3.8','Red'),('Healthcare workflow','5.0','Red'),('Innovation','5.6','Amber'),('Overall','5.2','Red')]
    table(d,['Dimension','Score / 10','Classification'],report,[4500,2200,2660],9)
    callout(d,'FINAL VERDICT: ⚠ NEEDS SIGNIFICANT WORK','Auriva is not ready for commercial launch. It is ready to earn evidence in a very narrow, assisted design-partner pilot only after every P0 recommendation is completed and independently verified.','FBE5E3')

def add_evidence_sheets(d):
    pagebreak(d); h(d, 'Appendix C — detailed launch evidence sheets', 1)
    para(d, 'The following sheets translate the headline audit into testable launch evidence. Each should be owned, rehearsed and signed before a controlled design-partner pilot. They are deliberately operational rather than aspirational.')
    sheets = [
        ('1. Account creation and clinic setup','An owner can create a clinic, configure safe defaults, invite the minimum team, add doctors/services and reach a ready-to-operate state without hidden support work.','The organization and membership model are real. A quantified first-value path, import journey and setup telemetry are not evidenced.','Measure time to ready-to-book, completion/drop-off by setup step, and support contacts per activation.','A non-technical clinic owner completes a rehearsal with no data ambiguity and receives a clear “ready” confirmation.'),
        ('2. Staff and patient identity','Staff access follows password/change rules; patients authenticate with OTP and select the correct family profile.','OTP design is thoughtful but production delivery is an open release gate. Identity unification remains future debt.','Record send, delivery, verify, expiry, abuse and failed-login metrics; test account role overlap.','Live OTP delivery works at target success rate; failures have safe recovery and support traceability.'),
        ('3. Online booking and preparation','A patient finds an appropriate clinician, sees reliable availability, submits required intake/consent, understands price/policy and receives confirmation.','Slot booking exists. Intake, consent, payment/deposit, delivery and waitlist are absent.','Track conversion by step, abandonment reason, confirmation delivery and no-show rate.','A pilot patient can book, prepare, cancel/reschedule and be reminded without calling the clinic.'),
        ('4. Walk-in identity and registration','Reception identifies a returning patient or creates a safe minimal profile in seconds.','Backend supports phone/name/DOB/Health ID, but core UI does not expose Health ID and family disambiguation is weak.','Time identity resolution, duplicate-profile rate and wrong-profile near misses.','Staff successfully identify family members using two identifiers and can see why a duplicate warning appeared.'),
        ('5. Reception queue control','Reception can see the room state, check in, send in, reassign, recover a wait breach and communicate changes.','The board clearly shows stage ownership and wait aging. Reassignment, capacity thresholds and outbound communication are missing.','Track wait percentile, queue exceptions, recovery time and manual workaround count.','A late doctor/absent patient/urgent walk-in scenario is resolved with authorized, audited actions.'),
        ('6. Doctor consultation','A doctor can safely find context, document, prescribe, order tests, save confidently and hand back to checkout.','The Workbench concept is strong. Clinical record is largely free text; autosave/conflict and clinical governance are insufficiently evidenced.','Measure consult task time, note completion, correction rate and clinician confidence.','Clinicians complete timed, realistic cases and can identify saved/signed/amended state without ambiguity.'),
        ('7. Prescription and diagnostics','A prescription is structured, signed, versioned, printable/deliverable and safely related to patient history; test results are reviewed and acknowledged.','Basic Rx and lab flows exist, but medication safety, robust result context and external delivery are absent.','Track Rx completion, amendment, interaction-warning handling and result acknowledgment.','Clinical Safety Board approves template/content and doctors can demonstrate safe correction/acknowledgment.'),
        ('8. Checkout and revenue closure','Service delivery becomes a transparent invoice, payment is captured in the right channel, a receipt is issued and follow-up is offered.','Invoice/payment/correction architecture is substantive. Online payment, active tax treatment and service-to-charge integration need completion.','Track completed-to-invoiced, invoiced-to-paid, concession/refund reason and outstanding aging.','A patient can pay in person or remotely, receive a receipt and be offered the next clinically appropriate booking.'),
        ('9. Patient after-visit experience','Patients receive what matters after care: documents, payment state, next action and a safe contact path.','The portal has warm navigation and records, but no outbound delivery, online payment or after-visit orchestration.','Track document delivery/open, payment completion, rebooking and care-plan engagement.','A patient who never reopens the app still receives consented essential information correctly.'),
        ('10. Owner decision system','An owner can see whether the clinic is flowing, collecting and retaining without building spreadsheets.','Command Center is a foundation; decision-grade KPIs and analytics are thin.','Track wait, no-show, collection, rebooking, plans and staff activity through a single dashboard.','The owner can answer “what needs attention today?” and “did we improve this month?” in under two minutes.'),
        ('11. Multi-role permissions','Every role sees only safe work, has a usable day-one surface and loses access immediately when changed.','RBAC/audit foundations are strong; nurse/technician action surfaces and clinic-scoped grants are incomplete.','Run permission matrix regression and role productivity tests.','No role is sold until its representative can complete the intended workflow end-to-end.'),
        ('12. Data integrity and privacy','Tenant boundaries, audit trails, backups, correction paths and patient sharing protect health data.','Many safeguards exist; restore evidence, formal privacy/compliance pack, family consent and scale hardening remain gaps.','Test cross-tenant denial, restore RTO/RPO, deletion/retention and audit completeness.','Security/privacy lead signs a release evidence pack and a restore drill passes.'),
        ('13. Accessibility and responsive use','Keyboard, screen reader, contrast, focus, motion and mobile behavior support real staff and patient diversity.','Component-level intent exists; independent conformance evidence does not.','Run automated and manual accessibility suite at desktop/tablet/mobile breakpoints.','Critical journeys complete with keyboard and screen reader; no critical contrast/touch defects remain.'),
        ('14. Performance and reliability','Morning opening, queue spikes, payments, booking and notifications perform predictably under expected load.','Health/ready endpoints and logging exist; unpaged lists, N+1 reads and single-instance mechanisms are documented risk.','Set SLOs for response time/error rate; load test realistic clinic concurrency.','SLOs pass with headroom and alert owners are on call for the pilot.'),
        ('15. Release quality gates','The product can be built, tested, deployed, monitored and rolled back reliably.','Tests/types pass, but lint fails and build relies on remote fonts in this environment.','Make lint/type/test/build, migration, backup, smoke and rollback evidence mandatory.','A release candidate clears all gates from a clean environment and is rehearsed through rollback.'),
        ('16. Customer success and support','A clinic knows what success means, where to get help and how feedback changes product decisions.','Release docs have operational notes; formal customer-success model is not evidenced.','Measure adoption, time-to-first-value, support response, workflow completion and NPS/CSAT.','Each pilot has named sponsor, success plan, escalation channel and weekly value review.'),
        ('17. Sales and packaging','A prospect understands fit, price, implementation, boundaries and proof without a bespoke founder explanation.','Edition structure is conceptual; price, migration and support offer are not ready.','Track qualified pipeline, conversion, cycle length, implementation time and early retention.','Sales uses a standard offer for one beachhead, with no capability promises outside shipped scope.'),
        ('18. Clinical governance and change control','Clinical workflow changes have accountable review, risk category, validation and release decision.','Product correctly defers AI diagnosis; a formal clinical governance operating system is not evidenced.','Audit change records, clinician review, safety incidents and exception approvals.','No clinical feature ships without documented review, usability validation and post-launch monitoring.')]
    for i,(title,purpose,gap,measure,exitcrit) in enumerate(sheets):
        pagebreak(d); h(d, title, 2)
        callout(d,'TEST INTENT',purpose,'E4EFEC')
        h(d,'Evidence observed',3); para(d, gap)
        h(d,'Measurement required',3); para(d, measure)
        h(d,'Pilot exit criterion',3); para(d, exitcrit)
        h(d,'Board challenge',3); para(d, 'Do not mark this sheet green because a component or API exists. Mark it green only when a representative user can complete the scenario, failure handling is clear, the result is auditable and the operating owner can explain the metric.')

def add_director_appendices(d):
    pagebreak(d); h(d,'Required leadership actions',1)
    h(d,'Product Office Recommendations',2)
    for x in ['Choose one beachhead and write a one-sentence value proposition with an observable 90-day outcome.', 'Move communications and pre-visit readiness ahead of broad feature expansion.', 'Use treatment planning only when supported by validated specialty demand.', 'Create a decision gate for every milestone: user evidence, success metric, commercial owner and stop condition.']:
        bullet(d,x)
    h(d,'CTO Recommendations',2)
    for x in ['Restore clean CI: lint/type/test/build; treat production build reproducibility as a release contract.', 'Close production configuration, backup/restore, monitoring, shared rate limit and scheduled-job risks.', 'Fund encounter data decomposition, search/pagination and multi-clinic correctness before Enterprise claims.', 'Publish an evidence-linked operational dashboard and incident runbook.']:
        bullet(d,x)
    h(d,'UX Director Recommendations',2)
    for x in ['Validate the queue board and Workbench in live clinic simulation before visual polish work.', 'Replace every visible false affordance/“Soon” feature with real capability, entitlement-aware hiding or a useful next action.', 'Institutionalize keyboard, touch, contrast, screen-reader and responsive acceptance tests.', 'Design the entire patient relationship: prepare → attend → pay → receive → follow up.']:
        bullet(d,x)
    h(d,'QA Director Recommendations',2)
    for x in ['Define P0 patient-safety/privacy/revenue-flow defect categories and stop-ship criteria.', 'Automate the full role-based day, including permission denial and failure/retry cases.', 'Add visual, accessibility, performance and production synthetic test layers.', 'Require a release evidence pack signed by Product, Engineering, Clinical Safety and Operations.']:
        bullet(d,x)
    h(d,'Go-To-Market Recommendations',2)
    for x in ['Do not sell “Healthcare OS” or Enterprise. Sell an assisted Professional pilot for one clinic type.', 'Recruit 3–5 design partners; baseline waiting time, no-shows, collection, rebooking and staff task time.', 'Build migration, implementation and support into the offer; do not treat them as afterthoughts.', 'Publish only claims that can be demonstrated in the product or measured in pilots.']:
        bullet(d,x)
    h(d,'Launch Checklist',2)
    checklist=[
        ('P0 product/clinical','OTP, communications, consent strategy, safety governance and E2E paths closed.'),
        ('P0 engineering','Lint/type/test/build clean; backups restored; alerts verified; production security configuration validated.'),
        ('P0 commercial','Beachhead, price, contract, onboarding, support owner and pilot KPIs approved.'),
        ('Pilot readiness','Design partners trained; migration dry-run; sandbox/rehearsal complete; escalation channel staffed.'),
        ('Launch evidence','One full clinic day completed twice without support intervention; no P0/P1 defects; decision review signed.')]
    table(d,['Gate','Exit criterion'],checklist,[2500,6860],9)
    h(d,'Final Launch Decision',2)
    para(d,'Do not authorize commercial launch. Authorize only a controlled, assisted design-partner pilot when P0 gates are green. Reassess after at least three clinics complete 30–90 days of measured usage, with documented evidence of activation, reliability, flow time, collection, rebooking, patient reachability and retention.')
    h(d,'Appendix A — current competitor sources',1)
    for x in [
        'Cliniko Online Payments: https://www.cliniko.com/features/transactions/online-payments/',
        'Jane Online Booking: https://jane.app/features/online-booking ; Jane Features: https://jane.app/us/features',
        'SimplePractice Client Portal: https://support.simplepractice.com/hc/en-us/articles/207925883 ; billing documents: https://support.simplepractice.com/hc/en-us/articles/360017844871-How-to-share-billing-documents-with-clients',
        'Halaxy Calendar: https://eu.halaxy.com/feature/calendar',
        'Practo Ray: https://www.practo.com/providers/clinics/ray',
        'DocPulse: https://docpulse.com/',
        'Clinicea Features: https://clinicea.com/features',
        'Healthie pricing/positioning: https://help.gethealthie.com/article/773-healthie-pricing']:
        bullet(d,x)
    h(d,'Appendix B — audit evidence references',1)
    for x in ['Release gates and known limitations: docs/RELEASE-CANDIDATE.md; production checklist: docs/production-checklist.md.', 'Workflow and UX gaps: docs/product-office-audit/01–05; current maturity: docs/knowledge-base/21-product-maturity-matrix.md.', 'Strategic scope and roadmap: docs/AURIVA-PRODUCT-BIBLE.md; docs/AURIVA-RELEASE-1X-ROADMAP.md.', 'Risk/debt: docs/technical-debt.md; docs/data-governance.md; prisma/schema.prisma; src/services and src/app/api.']:
        bullet(d,x)

def build():
    d=Document(); style_doc(d); add_title(d); add_method(d); pagebreak(d); add_exec(d); pagebreak(d); add_vision_market(d); pagebreak(d); add_journeys(d); pagebreak(d); add_business_quality(d); pagebreak(d); add_roadmap_launch(d); pagebreak(d); add_recs(d); pagebreak(d); add_screens(d); pagebreak(d); add_emotion_delight(d); add_evidence_sheets(d); add_director_appendices(d)
    d.save(DOCX)
    MD.write_text('# Auriva Independent Product Readiness Assessment\n\nThe formal report is available in `auriva-independent-product-readiness-assessment-2026-07-23.docx`. It contains the executive assessment, 100-item recommendation register, per-screen review, launch checklist and final launch decision.\n', encoding='utf-8')
    print(DOCX)

if __name__ == '__main__': build()
